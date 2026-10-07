"""Independent source readers and risk calculations; no production model imports."""
from __future__ import annotations

import math
from pathlib import Path
import json
import xml.etree.ElementTree as ET

import numpy as np
import pandas as pd

# Generic runtime/file budgets are shared; source-column declarations, currency
# checks, calendars, accounting and financial metrics remain independently implemented.
from .security_io import read_bounded_bytes, parse_bounded_xml, validate_xlsx, MAX_SHEET_COLUMNS
import openpyxl

RISK_FIELDS = ['ending_equity_usd', 'cagr', 'annualised_volatility', 'maximum_drawdown',
    'sharpe_ratio', 'beta_vs_unlevered_core', 'jensen_alpha', 'alpha_hac_se',
    'alpha_ci_low', 'alpha_ci_high', 'treynor_ratio', 'sortino_ratio', 'calmar_ratio']


def risk_statistics(equity, market, reference, capital):
    """Centered OLS and explicit lagged score products, ACT/360 and entry costs."""
    dates = equity.index
    if (not isinstance(dates, pd.DatetimeIndex) or not dates.equals(market.index)
        or len(dates) < 4 or dates.has_duplicates or not dates.is_monotonic_increasing
        or not math.isfinite(capital) or capital <= 0
        or not np.isfinite(equity).all() or not np.isfinite(market).all()
        or (equity <= 0).any() or (market <= 0).any()):
        raise ValueError('positive common-calendar wealth and capital required')
    rf = []
    for start, end in zip(dates[:-1], dates[1:]):
        rates = reference.reindex(pd.date_range(start, end - pd.Timedelta(days=1)))
        if not np.isfinite(rates).all():
            raise ValueError('complete reference intervals required')
        rf.append(math.fsum(map(float, rates)) / 360)
    r = equity.to_numpy()[1:] / equity.to_numpy()[:-1] - 1
    b = market.to_numpy()[1:] / market.to_numpy()[:-1] - 1
    r[0] = equity.iloc[1] / capital - 1
    b[0] = market.iloc[1] / capital - 1
    y, x = r - rf, b - rf
    centered = x - x.mean()
    if centered @ centered <= 0:
        raise ValueError('varying benchmark excess returns required')
    beta = (centered @ (y - y.mean())) / (centered @ centered)
    intercept = y.mean() - beta * x.mean()
    residual = y - intercept - beta * x
    design = np.column_stack([np.ones(len(x)), x])
    scores = design * residual[:, None]
    meat = np.zeros((2, 2))
    for a in range(2):
        for c in range(2):
            meat[a, c] = np.dot(scores[:, a], scores[:, c])
            for lag in range(1, min(5, len(x)-1)+1):
                meat[a, c] += (1-lag/6) * (np.dot(scores[lag:, a], scores[:-lag, c])
                    + np.dot(scores[:-lag, a], scores[lag:, c]))
    inverse = np.linalg.inv(design.T @ design)
    variance = (inverse @ meat @ inverse)[0, 0] * len(x)/(len(x)-2)
    se = math.sqrt(max(0., variance)) * 252
    alpha = intercept * 252
    values = equity.to_numpy()
    drawdown = values / np.maximum.accumulate(np.r_[capital, values])[1:] - 1
    cagr = (values[-1]/capital)**(365.2425/(dates[-1]-dates[0]).days)-1
    downside = math.sqrt(np.mean(np.minimum(y, 0)**2)*252)
    product_error = abs(math.prod(1+r)-values[-1]/capital)
    if product_error > 1e-11 * abs(values[-1]/capital):
        raise ValueError('investor returns do not compound to committed-capital wealth')
    return dict(ending_equity_usd=float(values[-1]), cagr=float(cagr),
        annualised_volatility=float(np.std(r, ddof=1)*math.sqrt(252)), maximum_drawdown=float(drawdown.min()),
        sharpe_ratio=float(np.mean(y)/np.std(y, ddof=1)*math.sqrt(252)),
        beta_vs_unlevered_core=float(beta), jensen_alpha=float(alpha), alpha_hac_se=float(se),
        alpha_ci_low=float(alpha-1.96*se), alpha_ci_high=float(alpha+1.96*se),
        treynor_ratio=float(y.mean()*252/beta) if beta > 0 else None,
        sortino_ratio=float(y.mean()*252/downside) if downside > 0 else None,
        calmar_ratio=float(cagr/abs(drawdown.min())) if drawdown.min() < 0 else None)


def spreadsheet_nav(path, sleeve):
    """Separate sparse-cell reader with independent worksheet/column declarations."""
    ns = '{urn:schemas-microsoft-com:office:spreadsheet}'
    root = parse_bounded_xml(read_bounded_bytes(path).replace(b'\xef\xbb\xbf', b''))
    sheet_name, column = ('Historical', 2) if sleeve == 'core' else ('Historical NAVs', 1)
    sheet = next((s for s in root.findall(ns+'Worksheet') if s.attrib.get(ns+'Name') == sheet_name), None)
    if sheet is None:
        raise ValueError('expected independent NAV worksheet missing')
    pairs = {}
    for row in sheet.iter(ns+'Row'):
        cells = {}; col = 0
        for cell in row.findall(ns+'Cell'):
            col = int(cell.attrib.get(ns+'Index', col+1))
            if not 1 <= col <= MAX_SHEET_COLUMNS:
                raise ValueError('Independent sparse column exceeds the declared bounds')
            data = cell.find(ns+'Data'); cells[col-1] = '' if data is None else data.text or ''
        try:
            day = pd.to_datetime(cells.get(0, '').replace('Sept', 'Sep').replace('-', '/'), format='%d/%b/%Y')
            value = float(cells.get(column, '').replace(',', ''))
        except (ValueError, TypeError):
            continue
        if day in pairs or not math.isfinite(value) or value <= 0:
            raise ValueError('invalid or duplicate independent NAV observation')
        if sleeve == 'core' and cells.get(1) != 'USD':
            raise ValueError('independent Core NAV currency differs')
        pairs[day] = value
    if not pairs:
        raise ValueError('no independent NAV observations')
    return pd.Series(pairs).sort_index()


def raw_source_checks(raw_root, spec, levels, reference):
    paths = {r['role']: Path(raw_root)/r['path'] for r in spec['sources']}
    checks = []
    def compare(role, actual, expected, limit):
        actual = actual.reindex(expected.index)
        if not np.isfinite(actual).all():
            raise ValueError('independent raw-source dates missing: '+role)
        error = float(np.max(np.abs(actual.to_numpy()-expected.to_numpy())))
        checks.append(dict(role=role, observations=len(expected), maximum_difference=error,
            allowed_difference=limit, passed=error <= limit))
    for sleeve in ['core', 'momentum', 'quality', 'value']:
        compare(sleeve, spreadsheet_nav(paths[sleeve], sleeve), levels[sleeve], 5e-7)
    validate_xlsx(paths['dimensional'])
    book = openpyxl.load_workbook(paths['dimensional'], read_only=True, data_only=False)
    try:
        sheet = book['Tabelle1']
        if sheet['B1'].value != spec['dimensional_security'] or sheet['B5'].value != 'USD':
            raise ValueError('independent Bloomberg identity/currency differs')
        pairs = {}
        for row in sheet.iter_rows(min_row=8, values_only=True):
            if isinstance(row[0], (int, float)) and isinstance(row[1], (int, float)):
                day = pd.Timestamp('1899-12-30')+pd.Timedelta(days=row[0])
                if day in pairs: raise ValueError('duplicate Bloomberg date')
                pairs[day] = row[1]
        compare('dimensional', pd.Series(pairs), levels.dimensional, 1e-10)
    finally:
        book.close()
    rates = {}
    validate_xlsx(paths['indicative_sofr'])
    book = openpyxl.load_workbook(paths['indicative_sofr'], read_only=True, data_only=True)
    try:
        for row in book['VWM Rates'].iter_rows(min_row=3, values_only=True):
            if isinstance(row[3], (int, float)) and row[0] is not None:
                day = pd.Timestamp(row[0])
                if day in rates: raise ValueError('duplicate indicative rate date')
                rates[day] = row[3]/10000
    finally:
        book.close()
    official = {}
    for record in json.loads(paths['official_sofr'].read_text())['refRates']:
        if record['type'] != 'SOFR': raise ValueError('independent reference identity differs')
        day = pd.Timestamp(record['effectiveDate'])
        if day in official: raise ValueError('duplicate official rate date')
        official[day] = record['percentRate']/100
    rates.update(official)
    series = pd.Series(rates).sort_index()
    carried = series.reindex(series.index.union(reference.index)).sort_index().ffill()
    compare('reference_rate_conversion_and_calendar_carry', carried, reference, 1e-10)
    if not all(c['passed'] for c in checks):
        raise ValueError('independent raw-source reader differs beyond preserved limits')
    return checks
