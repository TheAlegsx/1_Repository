"""Deterministic private-market-data build for spreads and leveraged references."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .security_io import validate_xlsx, read_xml_member, write_immutable_payloads, bounded_input
from openpyxl.utils.datetime import from_excel


DEFAULT_CUTOFF = date(2026, 8, 31)
RECENT_SPREAD_START = date(2025, 9, 1)
AMUNDI_ISIN = "FR0014010HV4"
LEVERAGED_INDEX_NAME = "MSCI World Leveraged 2X Daily (Net) USD Index"
XLSX_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


@dataclass(frozen=True)
class BloombergSpec:
    key: str
    security: str
    filename: str
    sleeve: str | None = None


@dataclass(frozen=True)
class BloombergSeries:
    spec: BloombergSpec
    data: pd.DataFrame
    source_sha256: str
    source_first_date: date
    source_last_date: date
    excluded_missing_px_last_dates: tuple[date, ...]
    excluded_incomplete_quote_dates: tuple[date, ...]


@dataclass(frozen=True)
class AmundiSeries:
    nav_usd: pd.Series
    source_sha256: str
    source_first_date: date
    source_last_date: date


@dataclass(frozen=True)
class MarketDataset:
    listing_quotes: pd.DataFrame
    spread_calibration: pd.DataFrame
    mxwoldnu_index: pd.DataFrame
    amundi_nav: pd.DataFrame
    benchmark_overlap: pd.DataFrame
    manifest: dict[str, object]


DEFAULT_BLOOMBERG_SPECS: tuple[BloombergSpec, ...] = (
    BloombergSpec("iwda", "IWDA LN Equity", "IWDA LN Equity grid_alspwryu.xlsx", "core"),
    BloombergSpec(
        "iwmo", "IWMO LN Equity", "IWMO LN Equity grid_udz0kbd5.xlsx", "momentum"
    ),
    BloombergSpec(
        "iwqu", "IWQU LN Equity", "IWQU LN Equity grid_xvx03ott.xlsx", "quality"
    ),
    BloombergSpec(
        "iwvl", "IWVL LN Equity", "IWVL LN Equity grid_ls5ij4wd.xlsx", "value"
    ),
    BloombergSpec(
        "mxwoldnu",
        "MXWOLDNU Index",
        "MXWOLDNU Index grid_rmi5kiny.xlsx",
    ),
)

DEFAULT_AMUNDI_FILENAME = (
    "NAV History_Amundi MSCI World (2x) Leveraged UCITS ETF Acc_"
    "FR0014010HV4_30_09_2025.xlsx"
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _number_or_nan(value: object) -> float:
    if value is None or isinstance(value, str):
        return math.nan
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"invalid numeric cell: {value!r}") from exc
    return result if math.isfinite(result) else math.nan


def _excel_date(value: object, epoch: datetime) -> pd.Timestamp:
    if isinstance(value, datetime):
        return pd.Timestamp(value.date())
    if isinstance(value, date):
        return pd.Timestamp(value)
    if isinstance(value, (int, float)):
        return pd.Timestamp(from_excel(value, epoch=epoch).date())
    raise ValueError(f"invalid Bloomberg date cell: {value!r}")


def read_bloomberg_hardcopy(
    path: str | Path,
    spec: BloombergSpec,
    *,
    cutoff: date = DEFAULT_CUTOFF,
) -> BloombergSeries:
    """Read the static sheet, never the Bloomberg add-in formula sheet."""

    source = Path(path)
    validate_xlsx(source)
    workbook = load_workbook(source, read_only=True, data_only=False)
    hardcopies = [sheet for sheet in workbook.worksheets if sheet.title != "Worksheet"]
    if len(hardcopies) != 1:
        raise ValueError(f"{source.name} must contain exactly one hard-copy sheet")
    sheet = hardcopies[0]

    if sheet["A1"].value != "Security" or sheet["B1"].value != spec.security:
        raise ValueError(f"unexpected security identity in {source.name}")
    if sheet["A5"].value != "Currency" or sheet["B5"].value != "USD":
        raise ValueError(f"{spec.security} is not labelled USD")
    headers = [sheet.cell(7, column).value for column in range(1, sheet.max_column + 1)]
    if not {"Date", "PX_LAST"}.issubset(headers):
        raise ValueError(f"{source.name} is missing Date or PX_LAST")
    positions = {str(value): index for index, value in enumerate(headers)}
    start_value = _excel_date(sheet["B2"].value, workbook.epoch).date()
    end_value = _excel_date(sheet["B3"].value, workbook.epoch).date()

    records: list[dict[str, object]] = []
    excluded_missing_px_last_dates: list[date] = []
    for values in sheet.iter_rows(min_row=8, values_only=True):
        if not values or values[positions["Date"]] is None:
            continue
        observation_date = _excel_date(values[positions["Date"]], workbook.epoch)
        if observation_date.date() > cutoff:
            continue
        raw_last = values[positions["PX_LAST"]]
        if raw_last is None or isinstance(raw_last, str):
            excluded_missing_px_last_dates.append(observation_date.date())
            continue
        record: dict[str, object] = {
            "date": observation_date,
            "px_last_usd": _number_or_nan(raw_last),
        }
        for field, output in (
            ("PX_BID", "px_bid_usd"),
            ("PX_ASK", "px_ask_usd"),
            ("VOLUME", "volume"),
        ):
            record[output] = (
                _number_or_nan(values[positions[field]])
                if field in positions and positions[field] < len(values)
                else math.nan
            )
        records.append(record)
    workbook.close()

    frame = pd.DataFrame.from_records(records).set_index("date").sort_index()
    frame.index = pd.DatetimeIndex(frame.index, name="date")
    if frame.empty or frame.index.has_duplicates or not frame.index.is_monotonic_increasing:
        raise ValueError(f"{spec.security} dates must be non-empty, unique, and sorted")
    if frame["px_last_usd"].isna().any() or (frame["px_last_usd"] <= 0.0).any():
        raise ValueError(f"{spec.security} contains invalid PX_LAST values")
    if end_value < cutoff:
        raise ValueError(f"{spec.security} source does not cover cutoff {cutoff}")
    if frame.index[-1].date() < cutoff - timedelta(days=4):
        raise ValueError(f"{spec.security} ends too far before cutoff {cutoff}")

    paired = frame[["px_bid_usd", "px_ask_usd"]].notna()
    incomplete_mask = paired.iloc[:, 0] != paired.iloc[:, 1]
    excluded_incomplete_quote_dates = tuple(
        value.date() for value in frame.index[incomplete_mask]
    )
    frame.loc[incomplete_mask, ["px_bid_usd", "px_ask_usd"]] = math.nan
    quoted = frame.dropna(subset=["px_bid_usd", "px_ask_usd"])
    if ((quoted["px_bid_usd"] <= 0.0) | (quoted["px_ask_usd"] <= 0.0)).any():
        raise ValueError(f"{spec.security} has non-positive bid/ask values")
    if (quoted["px_ask_usd"] < quoted["px_bid_usd"]).any():
        raise ValueError(f"{spec.security} has negative quoted spreads")
    midpoint = (quoted["px_ask_usd"] + quoted["px_bid_usd"]) / 2.0
    frame["full_spread_bps"] = (
        (frame["px_ask_usd"] - frame["px_bid_usd"]) / midpoint * 10_000.0
    )
    frame["half_spread_bps"] = frame["full_spread_bps"] / 2.0

    if frame.index[0].date() != start_value:
        raise ValueError(f"{spec.security} first observation differs from metadata")
    if frame.index[-1].date() > min(end_value, cutoff):
        raise ValueError(f"{spec.security} last observation exceeds metadata")
    return BloombergSeries(
        spec=spec,
        data=frame,
        source_sha256=_sha256(source),
        source_first_date=start_value,
        source_last_date=end_value,
        excluded_missing_px_last_dates=tuple(sorted(excluded_missing_px_last_dates)),
        excluded_incomplete_quote_dates=excluded_incomplete_quote_dates,
    )


def _xlsx_cells(path: Path) -> dict[str, str | float]:
    """Read values directly from OOXML so malformed style metadata is irrelevant."""

    validate_xlsx(path)
    with ZipFile(path) as archive:
        shared_root = read_xml_member(archive, "xl/sharedStrings.xml")
        shared = [
            "".join(node.text or "" for node in item.iter(f"{{{XLSX_NS}}}t"))
            for item in shared_root.findall(f"{{{XLSX_NS}}}si")
        ]
        sheet_root = read_xml_member(archive, "xl/worksheets/sheet1.xml")

    cells: dict[str, str | float] = {}
    for cell in sheet_root.iter(f"{{{XLSX_NS}}}c"):
        reference = cell.get("r")
        value_node = cell.find(f"{{{XLSX_NS}}}v")
        if reference is None or value_node is None or value_node.text is None:
            continue
        if cell.get("t") == "s":
            cells[reference] = shared[int(value_node.text)]
        else:
            cells[reference] = float(value_node.text)
    return cells


def read_amundi_nav(
    path: str | Path, *, cutoff: date = DEFAULT_CUTOFF
) -> AmundiSeries:
    source = Path(path)
    cells = _xlsx_cells(source)
    expected = {
        "C11": "Amundi MSCI World (2x) Leveraged UCITS ETF Acc",
        "C12": AMUNDI_ISIN,
        "C13": LEVERAGED_INDEX_NAME,
        "C14": "Synthetic",
        "C15": "USD",
        "B27": "Date",
        "C27": "Official NAV",
        "E27": "Nav Currency",
    }
    for reference, value in expected.items():
        if cells.get(reference) != value:
            raise ValueError(f"unexpected Amundi value at {reference}")

    records: list[tuple[pd.Timestamp, float]] = []
    source_dates: list[date] = []
    data_started = False
    for row in range(28, 100_000):
        raw_date = cells.get(f"B{row}")
        if raw_date is None:
            if data_started:
                break
            continue
        data_started = True
        currency = cells.get(f"E{row}")
        if currency != "USD":
            raise ValueError(f"unexpected Amundi NAV currency in row {row}")
        try:
            observation_date = pd.Timestamp(datetime.strptime(str(raw_date), "%d/%m/%Y"))
        except ValueError as exc:
            raise ValueError(f"invalid Amundi date in row {row}") from exc
        source_dates.append(observation_date.date())
        if observation_date.date() > cutoff:
            continue
        nav = _number_or_nan(cells.get(f"C{row}"))
        if not math.isfinite(nav) or nav <= 0.0:
            raise ValueError(f"invalid Amundi NAV in row {row}")
        records.append((observation_date, nav))

    series = pd.Series(
        [nav for _, nav in records],
        index=[observation_date for observation_date, _ in records],
        dtype=float,
        name="amundi_2x_nav_usd",
    ).sort_index()
    series.index = pd.DatetimeIndex(series.index, name="date")
    if series.empty or series.index.has_duplicates or not series.index.is_monotonic_increasing:
        raise ValueError("Amundi dates must be non-empty, unique, and sorted")
    if pd.Timestamp(cutoff) not in series.index:
        raise ValueError(f"Amundi has no observation on cutoff {cutoff}")
    return AmundiSeries(
        nav_usd=series,
        source_sha256=_sha256(source),
        source_first_date=series.index[0].date(),
        source_last_date=max(source_dates),
    )


def _spread_summary(
    listings: Iterable[BloombergSeries], cutoff: date
) -> pd.DataFrame:
    records: list[dict[str, object]] = []
    for listing in listings:
        if listing.spec.sleeve is None:
            continue
        for period, start in (
            ("full_history", listing.data.index[0].date()),
            ("recent_12m", RECENT_SPREAD_START),
        ):
            spreads = listing.data.loc[
                pd.Timestamp(start) : pd.Timestamp(cutoff), "full_spread_bps"
            ].dropna()
            if spreads.empty:
                raise ValueError(f"no {period} spread observations for {listing.spec.security}")
            records.append(
                {
                    "sleeve": listing.spec.sleeve,
                    "security": listing.spec.security,
                    "period": period,
                    "start_date": spreads.index[0].date().isoformat(),
                    "end_date": spreads.index[-1].date().isoformat(),
                    "observations": int(len(spreads)),
                    "full_spread_median_bps": spreads.median(),
                    "full_spread_p75_bps": spreads.quantile(0.75),
                    "full_spread_p95_bps": spreads.quantile(0.95),
                    "half_spread_median_bps": spreads.median() / 2.0,
                    "half_spread_p75_bps": spreads.quantile(0.75) / 2.0,
                    "half_spread_p95_bps": spreads.quantile(0.95) / 2.0,
                }
            )
    return pd.DataFrame.from_records(records)


def build_market_dataset(
    bloomberg_dir: str | Path,
    amundi_dir: str | Path,
    *,
    cutoff: date = DEFAULT_CUTOFF,
    bloomberg_specs: Iterable[BloombergSpec] = DEFAULT_BLOOMBERG_SPECS,
    amundi_filename: str = DEFAULT_AMUNDI_FILENAME,
) -> MarketDataset:
    """Build private listing evidence and leveraged-reference comparisons offline."""

    bloomberg_source = Path(bloomberg_dir)
    specs = tuple(bloomberg_specs)
    parsed = tuple(
        read_bloomberg_hardcopy(bloomberg_source / spec.filename, spec, cutoff=cutoff)
        for spec in specs
    )
    by_key = {item.spec.key: item for item in parsed}
    if "mxwoldnu" not in by_key:
        raise ValueError("Bloomberg specification is missing MXWOLDNU")
    amundi = read_amundi_nav(Path(amundi_dir) / amundi_filename, cutoff=cutoff)

    listing_frames: list[pd.DataFrame] = []
    for item in parsed:
        if item.spec.sleeve is None:
            continue
        frame = item.data.reset_index()
        frame.insert(1, "sleeve", item.spec.sleeve)
        frame.insert(2, "security", item.spec.security)
        listing_frames.append(frame)
    listing_quotes = pd.concat(listing_frames, ignore_index=True).sort_values(
        ["date", "sleeve"], ignore_index=True
    )
    spread_calibration = _spread_summary(parsed, cutoff)

    index_series = by_key["mxwoldnu"].data["px_last_usd"].rename(
        "mxwoldnu_index_level_usd"
    )
    mxwoldnu_index = index_series.to_frame()
    amundi_nav = amundi.nav_usd.to_frame()
    overlap = pd.concat([index_series, amundi.nav_usd], axis=1, join="inner").sort_index()
    overlap["mxwoldnu_daily_return"] = overlap["mxwoldnu_index_level_usd"].pct_change(
        fill_method=None
    )
    overlap["amundi_daily_return"] = overlap["amundi_2x_nav_usd"].pct_change(
        fill_method=None
    )
    overlap["amundi_minus_index_return"] = (
        overlap["amundi_daily_return"] - overlap["mxwoldnu_daily_return"]
    )
    comparison = overlap.dropna()
    if comparison.empty:
        raise ValueError("Amundi and MXWOLDNU have no return overlap")
    difference = comparison["amundi_minus_index_return"]

    source_entries = {
        item.spec.key: {
            "security": item.spec.security,
            "filename": item.spec.filename,
            "sha256": item.source_sha256,
            "source_first_date": item.source_first_date.isoformat(),
            "source_last_date": item.source_last_date.isoformat(),
            "admitted_last_date": item.data.index[-1].date().isoformat(),
            "admitted_observations": int(len(item.data)),
            "excluded_missing_px_last_dates": [
                value.isoformat() for value in item.excluded_missing_px_last_dates
            ],
            "excluded_incomplete_quote_dates": [
                value.isoformat() for value in item.excluded_incomplete_quote_dates
            ],
        }
        for item in parsed
    }
    source_entries["amundi_2x"] = {
        "product": "Amundi MSCI World (2x) Leveraged UCITS ETF Acc",
        "isin": AMUNDI_ISIN,
        "filename": amundi_filename,
        "sha256": amundi.source_sha256,
        "source_first_date": amundi.source_first_date.isoformat(),
        "source_last_date": amundi.source_last_date.isoformat(),
        "admitted_last_date": amundi.nav_usd.index[-1].date().isoformat(),
        "admitted_observations": int(len(amundi.nav_usd)),
    }
    manifest: dict[str, object] = {
        "schema_version": 1,
        "pipeline": "factor_portfolio.market_data",
        "network_required": False,
        "cutoff_date": cutoff.isoformat(),
        "currency": "USD",
        "release_status": "private vendor data and derived outputs; do not publish",
        "bloomberg_sheet_policy": "static hard-copy sheet only; live formula sheet ignored",
        "spread_definition": "(ask - bid) / midpoint * 10,000",
        "execution_cost_definition": "one-way half quoted spread relative to midpoint/NAV",
        "modeled_execution_cost_bps": 5.0,
        "modeled_execution_cost_decision": (
            "retain conservative common baseline; use 0 and 10 bp sensitivities"
        ),
        "recent_spread_window": {
            "start_date": RECENT_SPREAD_START.isoformat(),
            "end_date": cutoff.isoformat(),
        },
        "leveraged_index": LEVERAGED_INDEX_NAME,
        "benchmark_windows": {
            "mxwoldnu": {
                "start_date": mxwoldnu_index.index[0].date().isoformat(),
                "end_date": mxwoldnu_index.index[-1].date().isoformat(),
                "observations": int(len(mxwoldnu_index)),
            },
            "amundi_2x": {
                "start_date": amundi_nav.index[0].date().isoformat(),
                "end_date": amundi_nav.index[-1].date().isoformat(),
                "observations": int(len(amundi_nav)),
            },
            "common_overlap": {
                "start_date": overlap.index[0].date().isoformat(),
                "end_date": overlap.index[-1].date().isoformat(),
                "level_observations": int(len(overlap)),
                "return_observations": int(len(comparison)),
            },
        },
        "benchmark_comparison": {
            "daily_return_correlation": comparison[
                ["amundi_daily_return", "mxwoldnu_daily_return"]
            ].corr().iloc[0, 1],
            "mean_daily_difference_bps": difference.mean() * 10_000.0,
            "daily_mae_bps": difference.abs().mean() * 10_000.0,
            "daily_rmse_bps": float(np.sqrt(np.mean(np.square(difference)))) * 10_000.0,
            "amundi_cumulative_return": overlap["amundi_2x_nav_usd"].iloc[-1]
            / overlap["amundi_2x_nav_usd"].iloc[0]
            - 1.0,
            "mxwoldnu_cumulative_return": overlap["mxwoldnu_index_level_usd"].iloc[-1]
            / overlap["mxwoldnu_index_level_usd"].iloc[0]
            - 1.0,
        },
        "sources": source_entries,
    }
    return MarketDataset(
        listing_quotes=listing_quotes,
        spread_calibration=spread_calibration,
        mxwoldnu_index=mxwoldnu_index,
        amundi_nav=amundi_nav,
        benchmark_overlap=overlap,
        manifest=manifest,
    )


def write_market_dataset(
    dataset: MarketDataset, output_dir: str | Path
) -> dict[str, Path]:
    """Write byte-stable private outputs and refuse changed in-place snapshots."""

    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    paths = {
        "quotes": destination / "listing_quotes_usd.csv",
        "spreads": destination / "spread_calibration.csv",
        "index": destination / "mxwoldnu_index_usd.csv",
        "amundi": destination / "amundi_2x_nav_usd.csv",
        "overlap": destination / "leveraged_benchmark_overlap.csv",
        "manifest": destination / "market_data_manifest.json",
    }

    def csv_bytes(frame: pd.DataFrame, *, index: bool) -> bytes:
        buffer = io.StringIO(newline="")
        frame.to_csv(
            buffer,
            index=index,
            index_label="date" if index else None,
            date_format="%Y-%m-%d",
            float_format="%.12f",
            lineterminator="\n",
        )
        return buffer.getvalue().encode("utf-8")

    payloads = {
        paths["quotes"]: csv_bytes(dataset.listing_quotes, index=False),
        paths["spreads"]: csv_bytes(dataset.spread_calibration, index=False),
        paths["index"]: csv_bytes(dataset.mxwoldnu_index, index=True),
        paths["amundi"]: csv_bytes(dataset.amundi_nav, index=True),
        paths["overlap"]: csv_bytes(dataset.benchmark_overlap, index=True),
    }
    manifest = dict(dataset.manifest)
    manifest["canonical_files"] = {
        path.name: {"sha256": _sha256_bytes(payload), "rows": payload.count(b"\n") - 1}
        for path, payload in payloads.items()
    }
    payloads[paths["manifest"]] = (
        json.dumps(manifest, indent=2, sort_keys=True, ensure_ascii=False) + "\n"
    ).encode("utf-8")
    write_immutable_payloads(payloads, 'market-data')
    return paths


def load_market_dataset(input_dir: str | Path) -> MarketDataset:
    """Load private outputs only after checksum, row, and boundary verification."""

    source = Path(input_dir)
    manifest_path = source / "market_data_manifest.json"
    bounded_input(manifest_path)
    if not manifest_path.is_file():
        raise FileNotFoundError(manifest_path)
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as exc:
        raise ValueError("market-data manifest is not valid JSON") from exc

    file_map = {
        "quotes": "listing_quotes_usd.csv",
        "spreads": "spread_calibration.csv",
        "index": "mxwoldnu_index_usd.csv",
        "amundi": "amundi_2x_nav_usd.csv",
        "overlap": "leveraged_benchmark_overlap.csv",
    }
    records = manifest.get("canonical_files")
    if not isinstance(records, dict):
        raise ValueError("market-data manifest is missing canonical_files")
    for filename in file_map.values():
        record = records.get(filename)
        if not isinstance(record, dict) or not isinstance(record.get("sha256"), str):
            raise ValueError(f"market-data manifest is missing checksum for {filename}")
        path = source / filename
        bounded_input(path)
        if not path.is_file():
            raise FileNotFoundError(path)
        if _sha256(path) != record["sha256"]:
            raise ValueError(f"market-data checksum mismatch for {filename}")

    listing_quotes = pd.read_csv(source / file_map["quotes"], parse_dates=["date"])
    spread_calibration = pd.read_csv(source / file_map["spreads"])
    mxwoldnu_index = pd.read_csv(
        source / file_map["index"], index_col="date", parse_dates=["date"]
    )
    amundi_nav = pd.read_csv(
        source / file_map["amundi"], index_col="date", parse_dates=["date"]
    )
    benchmark_overlap = pd.read_csv(
        source / file_map["overlap"], index_col="date", parse_dates=["date"]
    )
    for frame in (mxwoldnu_index, amundi_nav, benchmark_overlap):
        frame.index = pd.DatetimeIndex(frame.index, name="date")

    frames = {
        file_map["quotes"]: listing_quotes,
        file_map["spreads"]: spread_calibration,
        file_map["index"]: mxwoldnu_index,
        file_map["amundi"]: amundi_nav,
        file_map["overlap"]: benchmark_overlap,
    }
    for filename, frame in frames.items():
        if len(frame) != records[filename].get("rows"):
            raise ValueError(f"market-data row count differs for {filename}")
    for label, frame in (
        ("MXWOLDNU", mxwoldnu_index),
        ("Amundi", amundi_nav),
        ("benchmark overlap", benchmark_overlap),
    ):
        if frame.empty or frame.index.has_duplicates or not frame.index.is_monotonic_increasing:
            raise ValueError(f"{label} dates must be non-empty, unique, and sorted")

    windows = manifest.get("benchmark_windows")
    if not isinstance(windows, dict):
        raise ValueError("market-data manifest is missing benchmark_windows")
    for key, frame in (
        ("mxwoldnu", mxwoldnu_index),
        ("amundi_2x", amundi_nav),
        ("common_overlap", benchmark_overlap),
    ):
        window = windows.get(key)
        if not isinstance(window, dict):
            raise ValueError(f"market-data manifest is missing {key} window")
        if frame.index[0].date().isoformat() != window.get("start_date"):
            raise ValueError(f"{key} start date differs from manifest")
        if frame.index[-1].date().isoformat() != window.get("end_date"):
            raise ValueError(f"{key} end date differs from manifest")

    return MarketDataset(
        listing_quotes=listing_quotes,
        spread_calibration=spread_calibration,
        mxwoldnu_index=mxwoldnu_index,
        amundi_nav=amundi_nav,
        benchmark_overlap=benchmark_overlap,
        manifest=manifest,
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Build private Bloomberg/Amundi evidence without network access."
    )
    parser.add_argument("--bloomberg-dir", type=Path, default=Path("data/raw/bloomberg"))
    parser.add_argument("--amundi-dir", type=Path, default=Path("data/raw/amundi"))
    parser.add_argument(
        "--output-dir", type=Path, default=Path("data/canonical/market_data")
    )
    parser.add_argument("--cutoff", type=date.fromisoformat, default=DEFAULT_CUTOFF)
    args = parser.parse_args(argv)

    dataset = build_market_dataset(
        args.bloomberg_dir, args.amundi_dir, cutoff=args.cutoff
    )
    paths = write_market_dataset(dataset, args.output_dir)
    print(
        f"Built private market evidence through {args.cutoff}: "
        f"{len(dataset.listing_quotes)} listing rows, "
        f"{len(dataset.benchmark_overlap)} benchmark-overlap rows."
    )
    for label, path in paths.items():
        print(f"{label}: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
